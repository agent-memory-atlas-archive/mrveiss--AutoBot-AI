# Copyright 2025-2026 mrveiss
# SPDX-License-Identifier: Apache-2.0
# AutoBot - AI-Powered Automation Platform
# Author: mrveiss
"""Heading/section structure on the canonical ExtractedDocument (#18826, owner option (a))."""

import io

import pytest

from media.document.extraction import ExtractedDocument, extract_document, extract_docx
from media.document.sections import Block, BlockKind, ListItem, Section, render_sections


def test_existing_producers_stay_compatible_sections_default_to_empty():
    assert ExtractedDocument(format="pdf", text="x").sections == ()
    assert extract_document(b"plain words").sections == ()


def test_render_sections_is_plain_text_with_tables_through_the_shared_renderer():
    sections = (
        Section(level=1, heading="Title", blocks=(Block(kind=BlockKind.PARAGRAPH, text="Intro."),)),
        Section(
            level=2,
            heading="Data",
            blocks=(
                Block(kind=BlockKind.TABLE, rows=(("Label", "Count"), ("bug", "578"))),
                Block(
                    kind=BlockKind.LIST,
                    items=(
                        ListItem("first", 0, True),
                        ListItem("nested", 1),
                        ListItem("second", 0, True),
                    ),
                ),
                Block(kind=BlockKind.CODE, text="  indented()\n"),
            ),
        ),
    )
    assert render_sections(sections) == (
        "Title\n\nIntro.\n\nData\n\nLabel | Count\nbug | 578\n\n1. first\n  - nested\n2. second\n\n  indented()\n"
    )


def test_extract_docx_text_is_unchanged_and_carries_no_sections_yet():
    """DOCX structure is #17985 (outline levels, not localised style names); text stays byte-identical."""
    pytest.importorskip("docx", reason="python-docx needed to synthesize a DOCX fixture")
    from docx import Document

    doc = Document()
    doc.add_heading("Chapter", level=1)
    doc.add_paragraph("Body one.")
    buffer = io.BytesIO()
    doc.save(buffer)

    extracted = extract_docx(buffer.getvalue())
    assert extracted.text == "Chapter\nBody one."
    assert extracted.sections == ()
